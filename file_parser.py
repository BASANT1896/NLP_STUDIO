"""Turn uploaded files into clean plain text."""
from __future__ import annotations

import io
import json
import re

import pandas as pd
import streamlit as st
from bs4 import BeautifulSoup
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pypdf import PdfReader

SUPPORTED = ["pdf", "docx", "pptx", "xlsx", "csv", "txt", "md", "json", "html", "htm"]
LEGACY = {"doc": ".docx", "ppt": ".pptx", "xls": ".xlsx"}


class ParseError(Exception):
    """Raised with a message that is safe to show to the user."""


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return data.decode("utf-8", errors="ignore")


def _pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ParseError("This PDF is password-protected. Remove the password and upload it again.")
    pages = []
    for page in reader.pages:
        t = (page.extract_text() or "").strip()
        if t:
            pages.append(t)
    return "\n\n".join(pages)


def _docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    blocks: list[str] = []
    for child in doc.element.body.iterchildren():  # keeps paragraphs and tables in reading order
        if child.tag.endswith("}p"):
            t = Paragraph(child, doc).text.strip()
            if t:
                blocks.append(t)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, doc).rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def _pptx(data: bytes) -> str:
    prs = Presentation(io.BytesIO(data))
    out = []
    for n, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"Notes: {notes}")
        if texts:
            out.append(f"Slide {n}\n" + "\n".join(texts))
    return "\n\n".join(out)


def _xlsx(data: bytes) -> str:
    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
    out = []
    for name, df in sheets.items():
        df = df.dropna(how="all")
        if not df.empty:
            out.append(f"Sheet: {name}\n{df.to_csv(index=False)}")
    return "\n\n".join(out)


def _json(data: bytes) -> str:
    raw = _decode(data)
    try:
        return json.dumps(json.loads(raw), indent=2, ensure_ascii=False)
    except ValueError:
        return raw


def _html(data: bytes) -> str:
    soup = BeautifulSoup(_decode(data), "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return soup.get_text("\n")


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


@st.cache_data(show_spinner="Reading file…", max_entries=8)
def extract_text(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in LEGACY:
        raise ParseError(f"Old .{ext} files aren't supported. Save the file as {LEGACY[ext]} and upload it again.")
    handlers = {
        "pdf": _pdf, "docx": _docx, "pptx": _pptx, "xlsx": _xlsx,
        "json": _json, "html": _html, "htm": _html,
        "csv": _decode, "txt": _decode, "md": _decode,
    }
    if ext not in handlers:
        raise ParseError(f"'.{ext}' files aren't supported. Try: {', '.join(SUPPORTED)}.")
    try:
        return _clean(handlers[ext](data))
    except ParseError:
        raise
    except Exception as e:  # noqa: BLE001
        raise ParseError(f"Couldn't read {filename}: {e}") from e
